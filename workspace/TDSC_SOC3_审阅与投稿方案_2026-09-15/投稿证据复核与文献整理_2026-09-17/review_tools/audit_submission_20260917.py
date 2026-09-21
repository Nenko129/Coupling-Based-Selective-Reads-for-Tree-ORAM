"""Read-only audit of the frozen paper bundle. Writes only to a separate review directory."""
from pathlib import Path
from collections import Counter, defaultdict
import json, csv, math, statistics, hashlib, zipfile, re
import xml.etree.ElementTree as ET

ROOT=Path('D:/projects/SDE-R0')
BASE=ROOT/'TDSC_SOC3_审阅与投稿方案_2026-09-15'
HOME=BASE/'Selective_通用组合与投稿实验准备_2026-09-16'
PACK=HOME/'论文实验章节完整证据_2026-09-16'
EVAL=BASE/'论文带宽实验_2026-09-16'
OUT=BASE/'投稿证据复核与文献整理_2026-09-17'
OUT.mkdir(exist_ok=True)
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def csvrows(name):
 with (PACK/'tables'/name).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
issues=[];counts=Counter();notes=[]
def ck(ok,tag,detail):
 counts[tag]+=1
 if not ok:issues.append({'check':tag,'detail':detail})
def close(a,b):return math.isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-7)
manifest=load(PACK/'evidence_manifest.json')
manifest_inventory=[]
for row in manifest['files']:
 p=PACK/row['copy'];ok=p.is_file() and sha(p)==row['sha256']
 ck(ok,'manifest_hash',row['copy'])
 if row.get('source'):
  s=Path(row['source']);ck(s.is_file() and sha(s)==row['sha256'],'source_copy_hash',row['source'])
 manifest_inventory.append((str(p),row['sha256']))
validation=load(PACK/'bundle_validation.json')
zipinfo=validation.get('zip',{})
zip_path=HOME/'TDSC论文实验阶段证据总包_2026-09-16.zip'
with zipfile.ZipFile(zip_path) as z:
 bad=z.testzip();ck(bad is None,'zip_crc',bad)
 zip_entries=len(z.namelist())
# Accommodate the frozen validation's explicit ZIP fields without guessing.
for key,v in validation.items():
 if 'zip' in key.lower() and 'sha' in key.lower():ck(sha(zip_path)==v,'zip_sha256',key)
 if isinstance(v,dict) and 'sha256' in v and 'zip' in key.lower():ck(sha(zip_path)==v['sha256'],'zip_sha256',key)
minimal={p.stem:load(p) for p in (PACK/'source_evidence/receipts_minimal_120').glob('*.json')}
scale={p.stem:load(p) for p in (PACK/'source_evidence/receipts_scale_125').glob('*.json')}
existing={str(p.relative_to(PACK)):load(p) for p in (PACK/'source_evidence/receipts_audited_existing').rglob('*.json')}
plan=load(PACK/'source_evidence/plans_and_locks/formal_plan.json')['specs']
sp=[r['spec'] for r in load(PACK/'source_evidence/plans_and_locks/extended_scale_queue.json')['items']]
for name,rs,ps in [('minimal',minimal,plan),('scale',scale,sp)]:
 expected={r['id']:r for r in ps};ck(set(rs)==set(expected),'exact_matrix_ids',name)
 for rid,r in rs.items():
  ck(r['status']=='passed','formal_status',rid)
  ck(r['spec']==expected[rid],'formal_spec',rid)

def audit_bill(b,label):
 if 'total_bytes' not in b or 'request_bytes' not in b:return
 ck(b['total_bytes']==b['request_bytes']+b['response_bytes'],'wire_direction_sum',label)
 if 'components' in b:ck(b['total_bytes']==sum(b['components'].values()),'component_sum',label)
 for key in ['by_stage','by_tree']:
  if key in b:ck(b['total_bytes']==sum(sum(x.values()) for x in b[key].values()),key+'_sum',label)
 if 'independently_invoiced_rpcs' in b:ck(b['rpc']==b['independently_invoiced_rpcs'],'rpc_invoice',label)
 if 'next_sequence' in b:ck(b['rpc']==b['next_sequence']-b['first_sequence'],'rpc_sequence',label)

raw_all=list(minimal.items())+list(scale.items())+list(existing.items())
status_counts=Counter()
source_refs=set();trace_refs=set();unbilled=[]
def hash_leaves(v):
 if isinstance(v,dict):
  for k,w in v.items():
   if isinstance(w,str) and re.fullmatch('[0-9a-f]{64}',w):yield k,w
   else:yield from hash_leaves(w)

for label,r in raw_all:
 status_counts[r.get('status','missing')]+=1
 source_refs.update(hash_leaves(r.get('source_hashes',{})))
 t=r.get('trace',{})
 if isinstance(t,dict) and t.get('path') and t.get('sha256'):trace_refs.add((t['path'],t['sha256']))
 for pn,p in r.get('phases',{}).items():
  if not isinstance(p,dict):continue
  if 'bills' in p:
   for i,b in enumerate(p['bills']):audit_bill(b,label+':'+pn+':'+str(i))
   for f in ['total_bytes','rpc']:
    if f in p:ck(p[f]==sum(b.get(f,0) for b in p['bills']),'backend_'+f+'_sum',label+':'+pn)
  else:audit_bill(p,label+':'+pn)
 m=r.get('phases',{}).get('measurement',{})
 if m and 'requests' in m and 'total_bytes' in m:
  ck(close(r['bytes_per_request'],m['total_bytes']/m['requests']),'bytes_per_request',label)
  ck(close(r['rpc_per_request'],m['rpc']/m['requests']),'rpc_per_request',label)
 else:unbilled.append(label)

def trace_id(r):
 t=r.get('trace',{});return t.get('sha256') if isinstance(t,dict) else str(t)
def paircheck(a,b,label):
 ck(bool(a.get('answer_sha256')) and a['answer_sha256']==b.get('answer_sha256'),'paired_answer',label)
 ck(bool(trace_id(a)) and trace_id(a)==trace_id(b),'paired_trace',label)
 if 'schedules' in a.get('phases',{}).get('measurement',{}):
  ck(a['phases']['measurement']['schedules']==b['phases']['measurement'].get('schedules'),'paired_frontend_schedule',label)
def statistic(xs):
 mean=statistics.mean(xs);half=2.7764451051977987*statistics.stdev(xs)/math.sqrt(len(xs))
 return mean,mean-half,mean+half
comparisons=[]
def compare(row,pairs,table,rpc_mode='paired'):
 vals=[100*(1-b['bytes_per_request']/a['bytes_per_request']) for a,b in pairs]
 obs=statistic(vals)
 for k,v in zip(['bytes_saving_pct','ci95_low_pct','ci95_high_pct'],obs):ck(close(row[k],v),'recomputed_'+k,table+':'+str(row))
 for k,v in [('baseline_bytes_per_request',statistics.mean(a['bytes_per_request'] for a,b in pairs)),('variant_bytes_per_request',statistics.mean(b['bytes_per_request'] for a,b in pairs))]:
  if k in row:ck(close(row[k],v),'recomputed_'+k,table)
 if row.get('rpc_saving_pct'):
  rv=statistics.mean(100*(1-b['rpc_per_request']/a['rpc_per_request']) for a,b in pairs) if rpc_mode=='paired' else 100*(1-statistics.mean(b['rpc_per_request'] for a,b in pairs)/statistics.mean(a['rpc_per_request'] for a,b in pairs))
  ck(close(row['rpc_saving_pct'],rv),'recomputed_rpc_saving_pct',table)
 ck(len(pairs)==5,'paired_n',table)
 for a,b in pairs:paircheck(a,b,table+':'+a['spec']['id'])
 comparisons.append({'table':table,'key':{k:row[k] for k in ['comparison','family','layout','B','N','workload','baseline','step'] if k in row},'n':len(pairs),'mean':obs[0],'ci95':[obs[1],obs[2]],'per_seed':vals})

mi={(r['spec']['family'],r['spec']['layout'],r['spec']['B'],r['spec']['arm'],r['spec']['trace_seed']):r for r in minimal.values()}
for name in ['table_03_minimal_composition.csv','table_04_attribution_controls.csv']:
 for row in csvrows(name):
  pre=(row['family'],row['layout'],int(row['B']))
  compare(row,[(mi[pre+(row['baseline'],s)],mi[pre+('selective',s)]) for s in range(101,106)],name,'ratio_of_means')
si={(r['spec']['N'],r['spec']['B'],r['spec']['kind'],r['spec']['trace_seed']):r for r in scale.values()}
for row in csvrows('table_02_scale.csv'):
 pre=(int(row['N']),int(row['B']))
 compare(row,[(si[pre+(row['baseline'],s)],si[pre+(row['variant'],s)]) for s in range(101,106)],'table_02_scale.csv')
tuned=[r for p,r in existing.items() if '/formal_tuned/' in p.replace('\\','/')]
ti={(r['spec']['kind'],r['spec']['workload'],r['spec']['trace_seed']):r for r in tuned}
for row in csvrows('table_01_tuned_main.csv'):
 compare(row,[(ti[(row['baseline'],row['workload'],s)],ti[(row['variant'],row['workload'],s)]) for s in range(101,106)],'table_01_tuned_main.csv')

# Reconstruct the ablation's named aliases from the audit's raw table, then verify
# each aliased observation against its physical receipt (never treat aliases as runs).
ab=load(PACK/'source_evidence/audits/eval__ablation_complete_audit.json')
audit_ablation_schema={'raw_type':type(ab.get('raw')).__name__,'raw_first':str(ab.get('raw',[])[:1])[:1600], 'receipts_first':str(ab.get('receipts',[])[:1])[:1000]}
physical={r['spec']['id']:r for r in existing.values()}
aliases={r['id']:r for r in ab['raw']}
for a in ab['raw']:
 p=physical[a['physical_id']]
 for k in ['bytes_per_request','rpc_per_request']:
  ck(close(a[k],p[k]),'ablation_alias_physical_value',a['id']+':'+k)
for row in csvrows('table_05_ablation.csv'):
 pairs=[]
 for s in range(101,106):
  a=aliases[f"B2_R0_{row['before']}_seed{s}"];b=aliases[f"B2_R0_{row['after']}_seed{s}"]
  pairs.append((physical[a['physical_id']],physical[b['physical_id']]))
 compare(row,pairs,'table_05_ablation.csv')

# Recorded sources can refer to historical versions; preserve every mismatch in the report.
resolved_sources=[];unresolved_sources=[];mismatched_sources=[]
for path,h in sorted(source_refs):
 candidates=[ROOT/path,EVAL/path,HOME/path,EVAL/'src'/path,HOME/'src'/path,HOME/'tools'/path]
 found=next((p for p in candidates if p.is_file()),None)
 if found is None:unresolved_sources.append({'path':path,'sha256':h})
 elif sha(found)!=h:mismatched_sources.append({'path':str(found),'expected':h,'current':sha(found)})
 else:resolved_sources.append(str(found))
trace_checks=[]
# Resolve descriptive hash keys (e.g. "driver") by exact digest in project sources.
by_digest=defaultdict(list)
for root in [BASE,ROOT/'SOC3_audit_2026-09-15']:
 for p in root.rglob('*.py'):
  if '__pycache__' in p.parts:continue
  by_digest[sha(p)].append(str(p))
alias_sources=[];still_unresolved=[]
for row in unresolved_sources:
 found=by_digest.get(row['sha256'],[])
 if found:alias_sources.append({**row,'matching_files':found})
 else:still_unresolved.append(row)
unresolved_sources=still_unresolved
for path,h in sorted(trace_refs):
 found=next((p for p in [Path(path),ROOT/path,EVAL/path,HOME/path] if p.is_file()),None)
 ok=bool(found and sha(found)==h);ck(ok,'trace_file_hash',path);trace_checks.append({'path':path,'ok':ok})

# Workbook audit is read-only: inspect saved OOXML, formulas, cached values and chart refs.
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
workbook=PACK/'TDSC论文实验数据总表.xlsx';sheets=[];formula_missing=[];errors=[];chart_refs=[]
with zipfile.ZipFile(workbook) as z:
 wb=ET.fromstring(z.read('xl/workbook.xml'));sheet_names=[x.get('name') for x in wb.find('m:sheets',ns)]
 for i,name in enumerate(sheet_names,1):
  node=ET.fromstring(z.read(f'xl/worksheets/sheet{i}.xml'));cs=node.findall('.//m:c',ns);fs=[c for c in cs if c.find('m:f',ns) is not None]
  for c in cs:
   if c.get('t')=='e':errors.append({'sheet':name,'cell':c.get('r'),'error':c.findtext('m:v',namespaces=ns)})
  for c in fs:
   if c.find('m:v',ns) is None or c.findtext('m:v',namespaces=ns) in [None,'']:formula_missing.append({'sheet':name,'cell':c.get('r')})
  dim=node.find('m:dimension',ns)
  sheets.append({'name':name,'stored_cells':len(cs),'formulas':len(fs),'dimension':dim.get('ref') if dim is not None else None})
 for name in z.namelist():
  if '/charts/' in name and name.endswith('.xml') and '/_rels/' not in name:
   node=ET.fromstring(z.read(name));refs=[x.text for x in node.findall('.//c:f',ns)];chart_refs.append({'chart':name,'references':refs})
   ck(all('#REF!' not in (x or '') for x in refs),'chart_valid_reference',name)
ck(not errors,'xlsx_error_cells',errors);ck(not formula_missing,'xlsx_formula_cache',formula_missing)
data=load(PACK/'workbook_data.json')
import openpyxl
wb_values=openpyxl.load_workbook(workbook,read_only=False,data_only=True)
xlsx_mapping=[
 ('主比较','tuned',4,['comparison','workload','baseline','variant','n','baseline_bytes_per_request','variant_bytes_per_request','bytes_saving_pct','ci95_low_pct','ci95_high_pct','rpc_saving_pct','scope']),
 ('规模矩阵','scale',4,['comparison','N','B','baseline','variant','n','baseline_bytes_per_request','variant_bytes_per_request','bytes_saving_pct','ci95_low_pct','ci95_high_pct','rpc_saving_pct']),
 ('最小组合','minimal_primary',4,['family','layout','B','baseline','variant','n','baseline_bytes_per_request','variant_bytes_per_request','bytes_saving_pct','ci95_low_pct','ci95_high_pct','theory_saving_pct','theory_gap_pp','rpc_saving_pct']),
 ('消融','ablation',4,['step','before','after','n','bytes_saving_pct','ci95_low_pct','ci95_high_pct','rpc_saving_pct','note'])]
for sheet,key,header,keys in xlsx_mapping:
 for i,row in enumerate(data[key],header+1):
  for j,k in enumerate(keys,1):
   a=wb_values[sheet].cell(i,j).value;b=row.get(k)
   ck(close(a,b) if isinstance(b,(int,float)) and not isinstance(b,bool) else a==b,'xlsx_saved_value',f'{sheet}!{i},{j}:{k}')
expected_summary=[data['tuned'][0]['bytes_saving_pct'],data['tuned'][1]['bytes_saving_pct'],data['tuned'][2]['bytes_saving_pct'],min(r['bytes_saving_pct'] for r in data['minimal_primary']),max(r['bytes_saving_pct'] for r in data['minimal_primary']),245]
for i,v in enumerate(expected_summary,6):ck(close(wb_values['总览'].cell(i,2).value,v),'xlsx_summary_formula_cache',f'B{i}')
wb_values.close()
mapping={'table_01_tuned_main.csv':'tuned','table_02_scale.csv':'scale','table_03_minimal_composition.csv':'minimal_primary','table_04_attribution_controls.csv':'minimal_controls','table_05_ablation.csv':'ablation'}
for f,key in mapping.items():
 rows=csvrows(f);vals=data[key];ck(len(rows)==len(vals),'workbook_data_rowcount',f)
 for i,(a,b) in enumerate(zip(rows,vals)):
  for k,v in a.items():
   if k not in b:continue
   target=b[k]
   eq=(not v and target is None) or (isinstance(target,(int,float)) and not isinstance(target,bool) and close(v,target)) or (str(target)==v)
   ck(eq,'workbook_data_csv',f+':'+str(i)+':'+k)

result={'checked_on':'2026-09-17','method':'read-only recomputation from saved receipts; not a fresh rerun or independent third-party audit',
 'status':'passed_with_scope_notes' if not issues else 'issues_found','counts':dict(counts),'issues':issues,
 'manifest_entries':len(manifest['files']),'zip_entries':zip_entries,'zip_sha256':sha(zip_path),
 'receipts':{'minimal':len(minimal),'scale':len(scale),'existing':len(existing),'statuses':dict(status_counts),'unbilled':unbilled},
 'comparisons':comparisons,'ablation_schema':audit_ablation_schema,
 'sources':{'resolved_matches':len(resolved_sources),'alias_resolved':alias_sources,'unresolved':unresolved_sources,'current_hash_mismatches':mismatched_sources},
 'trace_files':trace_checks,'workbook':{'sha256':sha(workbook),'sheets':sheets,'errors':errors,'missing_formula_cache':formula_missing,'charts':chart_refs},
 'scope_limits':['Fixed-window core/scale results; not stationary estimates','Minimal base->selective also changes placement; full_probe controls attribution','Frontend leakage/clock contract applies to Freecursive/rho','Capacity certificate is profile specific','No production network latency or true trusted peak memory measurement','No native four-author-system reproduction']}
(OUT/'evidence_audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k not in ['comparisons','trace_files','workbook','sources']},ensure_ascii=False)[:8500])
print('SOURCE_SUMMARY',len(resolved_sources),'unresolved',len(unresolved_sources),'mismatches',len(mismatched_sources))
print('SOURCE_DETAIL',json.dumps({'unresolved':unresolved_sources,'mismatches':mismatched_sources},ensure_ascii=False)[:4000])
print('WORKBOOK',json.dumps(result['workbook'],ensure_ascii=False)[:2000])
